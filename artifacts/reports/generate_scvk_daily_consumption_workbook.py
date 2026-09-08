from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import bindparam, text

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.db.connect import ENGINE_PG


REPORT_IDENTIFIERS = (
    "SCVK_B1",
    "SCVK_DP",
    "SCVK_DV",
    "SCVK_GR",
    "SCVK_HE",
    "SCVK_S1",
)

REPORT_START = date(2025, 7, 1)
REPORT_END_INCLUSIVE = date(2026, 6, 30)
OUTPUT_PATH = Path("artifacts/reports/scvk_denni_spotreba_2025-07-01_2026-06-30.xlsx")


def iter_days(start: date, end_inclusive: date) -> list[date]:
    days: list[date] = []
    cursor = start
    while cursor <= end_inclusive:
        days.append(cursor)
        cursor += timedelta(days=1)
    return days


def to_float(value: object) -> float | None:
    if value is None:
        return None
    return round(float(value), 3)


def load_daily_consumption() -> dict[tuple[str, date], dict[str, object]]:
    statement = text(
        """
        SELECT
            identifikace,
            date::date AS day,
            COUNT(*) AS measurement_count,
            COUNT(*) FILTER (WHERE COALESCE(reset_detected, FALSE)) AS reset_count,
            SUM(delta) FILTER (
                WHERE delta IS NOT NULL
                  AND delta >= 0
                  AND COALESCE(reset_detected, FALSE) = FALSE
            ) AS consumption
        FROM monitoring."Mereni_vodomery_vse"
        WHERE identifikace IN :identifiers
          AND date >= :period_start
          AND date < :period_end
          AND platne = TRUE
        GROUP BY identifikace, date::date
        """
    ).bindparams(bindparam("identifiers", expanding=True))

    result: dict[tuple[str, date], dict[str, object]] = {}
    with ENGINE_PG.connect() as conn:
        rows = conn.execute(
            statement,
            {
                "identifiers": list(REPORT_IDENTIFIERS),
                "period_start": datetime.combine(REPORT_START, datetime.min.time()),
                "period_end": datetime.combine(
                    REPORT_END_INCLUSIVE + timedelta(days=1),
                    datetime.min.time(),
                ),
            },
        ).mappings().all()

    for row in rows:
        result[(str(row["identifikace"]), row["day"])] = {
            "measurement_count": int(row["measurement_count"] or 0),
            "reset_count": int(row["reset_count"] or 0),
            "consumption": to_float(row["consumption"]),
        }
    return result


def autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 32)


def style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    autosize_columns(worksheet)


def write_workbook() -> None:
    days = iter_days(REPORT_START, REPORT_END_INCLUSIVE)
    daily_consumption = load_daily_consumption()

    workbook = Workbook()
    workbook.remove(workbook.active)

    for identifier in REPORT_IDENTIFIERS:
        worksheet = workbook.create_sheet(identifier)
        worksheet.append(["Datum", "Spotreba m3", "Pocet mereni", "Pocet resetu", "Stav dat"])

        for day in days:
            aggregate = daily_consumption.get((identifier, day), {})
            measurement_count = int(aggregate.get("measurement_count") or 0)
            reset_count = int(aggregate.get("reset_count") or 0)
            consumption = aggregate.get("consumption")
            if measurement_count == 0:
                status = "BEZ_MERENI"
            elif consumption is None:
                status = "CHYBI_DELTA"
            elif reset_count:
                status = "OK_RESET_V_OBDOBI"
            else:
                status = "OK"
            worksheet.append([day, consumption, measurement_count, reset_count, status])

        for row in worksheet.iter_rows(min_row=2):
            row[0].number_format = "yyyy-mm-dd"
            if isinstance(row[1].value, float):
                row[1].number_format = "0.000"
        style_sheet(worksheet)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)


def main() -> None:
    write_workbook()
    print(f"Vytvoreno: {OUTPUT_PATH.resolve()}")
    print(f"Pocet zalozek: {len(REPORT_IDENTIFIERS)}")
    print(f"Radku na zalozku bez hlavicky: {len(iter_days(REPORT_START, REPORT_END_INCLUSIVE))}")


if __name__ == "__main__":
    main()
