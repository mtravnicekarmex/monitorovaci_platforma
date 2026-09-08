from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import bindparam, text

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.db.connect import ENGINE_PG


REPORT_IDENTIFIERS = (
    "SCVK_B1",
    "SCVK_DP",
    "SCVK_DV",
    "SCVK_GR",
    "SCVK_HE",
    "SCVK_S1",
)

METER_METADATA = {
    "SCVK_B1": {
        "name": "B1 NEWAYS",
        "supply_point": "681189817",
        "serial": "015971",
    },
    "SCVK_DP": {
        "name": "DOKTOR pozarni voda",
        "supply_point": "602003779",
        "serial": "040202",
    },
    "SCVK_DV": {
        "name": "DOKTOR voda",
        "supply_point": "602003694",
        "serial": "017021",
    },
    "SCVK_GR": {
        "name": "GROBAR",
        "supply_point": "602009738",
        "serial": "005354",
    },
    "SCVK_HE": {
        "name": "HECHT",
        "supply_point": "681041132",
        "serial": "005872",
    },
    "SCVK_S1": {
        "name": "S1",
        "supply_point": "681290535",
        "serial": "013445",
    },
}

REPORT_START = datetime(2025, 7, 1)
REPORT_END = datetime(2026, 8, 1)
OUTPUT_PATH = Path("artifacts/reports/scvk_mesicni_spotreba_2025-07_2026-07.xlsx")


@dataclass(frozen=True)
class MonthPeriod:
    label: str
    start: datetime
    end: datetime


def iter_months(start: datetime, exclusive_end: datetime) -> list[MonthPeriod]:
    periods: list[MonthPeriod] = []
    cursor = start
    while cursor < exclusive_end:
        if cursor.month == 12:
            next_month = datetime(cursor.year + 1, 1, 1)
        else:
            next_month = datetime(cursor.year, cursor.month + 1, 1)
        periods.append(
            MonthPeriod(
                label=f"{cursor.year}-{cursor.month:02d}",
                start=cursor,
                end=next_month,
            )
        )
        cursor = next_month
    return periods


def to_float(value: object) -> float | None:
    if value is None:
        return None
    return round(float(value), 3)


def load_last_values_before(
    cutoffs: list[datetime],
) -> dict[datetime, dict[str, tuple[float, datetime]]]:
    statement = text(
        """
        WITH ranked_measurements AS (
            SELECT
                identifikace,
                objem,
                date,
                ROW_NUMBER() OVER (
                    PARTITION BY identifikace
                    ORDER BY date DESC, id DESC
                ) AS row_num
            FROM monitoring."Mereni_vodomery_vse"
            WHERE identifikace IN :identifiers
              AND date < :cutoff
              AND platne = TRUE
              AND objem IS NOT NULL
        )
        SELECT identifikace, objem, date
        FROM ranked_measurements
        WHERE row_num = 1
        """
    ).bindparams(bindparam("identifiers", expanding=True))

    values_by_cutoff: dict[datetime, dict[str, tuple[float, datetime]]] = {}
    with ENGINE_PG.connect() as conn:
        for cutoff in cutoffs:
            rows = conn.execute(
                statement,
                {"identifiers": list(REPORT_IDENTIFIERS), "cutoff": cutoff},
            ).all()
            values_by_cutoff[cutoff] = {
                str(identifier): (to_float(volume), timestamp)
                for identifier, volume, timestamp in rows
                if to_float(volume) is not None
            }
    return values_by_cutoff


def load_monthly_consumption() -> dict[tuple[str, str], dict[str, object]]:
    statement = text(
        """
        SELECT
            identifikace,
            COUNT(*) AS measurement_count,
            COUNT(*) FILTER (WHERE COALESCE(reset_detected, FALSE)) AS reset_count,
            SUM(delta) FILTER (
                WHERE delta IS NOT NULL
                  AND delta >= 0
                  AND COALESCE(reset_detected, FALSE) = FALSE
            ) AS consumption
        FROM monitoring."Mereni_vodomery_vse"
        WHERE identifikace IN :identifiers
          AND date >= :period_start
          AND date < :period_end
          AND platne = TRUE
        GROUP BY identifikace
        """
    ).bindparams(bindparam("identifiers", expanding=True))

    result: dict[tuple[str, str], dict[str, object]] = {}
    with ENGINE_PG.connect() as conn:
        for month in iter_months(REPORT_START, REPORT_END):
            rows = conn.execute(
                statement,
                {
                    "identifiers": list(REPORT_IDENTIFIERS),
                    "period_start": month.start,
                    "period_end": month.end,
                },
            ).mappings().all()
            for row in rows:
                result[(month.label, str(row["identifikace"]))] = {
                    "measurement_count": int(row["measurement_count"] or 0),
                    "reset_count": int(row["reset_count"] or 0),
                    "consumption": to_float(row["consumption"]),
                }
    return result


def build_rows() -> list[dict[str, object]]:
    months = iter_months(REPORT_START, REPORT_END)
    cutoffs = sorted({month.start for month in months} | {month.end for month in months})
    snapshot_values = load_last_values_before(cutoffs)
    monthly_consumption = load_monthly_consumption()

    rows: list[dict[str, object]] = []
    for month in months:
        for identifier in REPORT_IDENTIFIERS:
            metadata = METER_METADATA[identifier]
            start_snapshot = snapshot_values.get(month.start, {}).get(identifier)
            end_snapshot = snapshot_values.get(month.end, {}).get(identifier)
            start_value = start_snapshot[0] if start_snapshot else None
            end_value = end_snapshot[0] if end_snapshot else None
            aggregate = monthly_consumption.get((month.label, identifier), {})
            measurement_count = int(aggregate.get("measurement_count") or 0)
            reset_count = int(aggregate.get("reset_count") or 0)
            consumption = aggregate.get("consumption")

            state_consumption = None
            state_note = "OK"
            if start_value is None or end_value is None:
                state_note = "HRANICNI_STAV_NEUPLNY"
            elif end_value < start_value:
                state_note = "RESET_NEBO_VYMENA_VODOMERU"
            else:
                state_consumption = round(end_value - start_value, 3)

            if measurement_count == 0:
                data_status = "BEZ_MERENI"
            elif consumption is None:
                data_status = "CHYBI_DELTA"
            elif reset_count:
                data_status = "OK_RESET_V_OBDOBI"
            else:
                data_status = "OK"

            rows.append(
                {
                    "Obdobi": month.label,
                    "Identifikace": identifier,
                    "Odberne misto": metadata["supply_point"],
                    "Nazev": metadata["name"],
                    "Cislo vodomeru": metadata["serial"],
                    "Pocatecni stav m3": start_value,
                    "Datum pocatecniho stavu": start_snapshot[1] if start_snapshot else None,
                    "Konecny stav m3": end_value,
                    "Datum konecneho stavu": end_snapshot[1] if end_snapshot else None,
                    "Spotreba ze stavu m3": state_consumption,
                    "Spotreba m3": consumption,
                    "Pocet mereni": measurement_count,
                    "Pocet resetu": reset_count,
                    "Stav hranicnich stavu": state_note,
                    "Stav dat": data_status,
                }
            )
    return rows


def autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 42)


def style_table(worksheet, header_row: int = 1) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)
    autosize_columns(worksheet)


def append_rows(worksheet, rows: list[dict[str, object]], headers: list[str]) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])


def write_workbook(rows: list[dict[str, object]]) -> None:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Detail"
    detail_headers = [
        "Obdobi",
        "Identifikace",
        "Odberne misto",
        "Nazev",
        "Cislo vodomeru",
        "Pocatecni stav m3",
        "Datum pocatecniho stavu",
        "Konecny stav m3",
        "Datum konecneho stavu",
        "Spotreba ze stavu m3",
        "Spotreba m3",
        "Pocet mereni",
        "Pocet resetu",
        "Stav hranicnich stavu",
        "Stav dat",
    ]
    append_rows(detail, rows, detail_headers)
    style_table(detail)

    overview = workbook.create_sheet("Prehled")
    overview.append(["Obdobi", *REPORT_IDENTIFIERS])
    by_month_identifier = {
        (str(row["Obdobi"]), str(row["Identifikace"])): row["Spotreba m3"]
        for row in rows
    }
    for month in iter_months(REPORT_START, REPORT_END):
        overview.append(
            [
                month.label,
                *[
                    by_month_identifier.get((month.label, identifier))
                    for identifier in REPORT_IDENTIFIERS
                ],
            ]
        )
    style_table(overview)

    metadata = workbook.create_sheet("Metadata")
    metadata.append(["Polozka", "Hodnota"])
    metadata.append(["Obdobi od", "2025-07"])
    metadata.append(["Obdobi do", "2026-07"])
    metadata.append(["Zdroj", 'monitoring."Mereni_vodomery_vse"'])
    metadata.append(["Filtr platnosti", "platne = TRUE, objem IS NOT NULL"])
    metadata.append(["Metoda spotreby", "soucet platnych nezapornych intervalovych delta hodnot v mesici"])
    metadata.append(["Kontrola hranic", "detail obsahuje rozdil poslednich platnych stavu pred zacatkem a koncem mesice"])
    metadata.append(["Vygenerovano", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    style_table(metadata)

    for worksheet in (detail, overview):
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)


def main() -> None:
    rows = build_rows()
    write_workbook(rows)
    ok_rows = sum(1 for row in rows if str(row["Stav dat"]).startswith("OK"))
    print(f"Vytvoreno: {OUTPUT_PATH.resolve()}")
    print(f"Radku detailu: {len(rows)}")
    print(f"Radku se stavem OK: {ok_rows}")


if __name__ == "__main__":
    main()
